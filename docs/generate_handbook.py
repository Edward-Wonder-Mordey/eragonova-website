from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

OUT = 'docs/Eragonova_Enterprise_Website_Owners_Handbook.docx'
BADGE = 'docs/.eragonova_handbook_badge.png'

NAVY='0F172A'; BLUE='2563EB'; CYAN='0EA5E9'; INDIGO='4F46E5'; SLATE='334155'; MUTED='64748B'
LIGHT='F8FAFC'; LIGHTBLUE='EFF6FF'; WHITE='FFFFFF'; GREEN='ECFDF5'; AMBER='FFFBEB'; RED='FEF2F2'; BORDER='CBD5E1'; TEAL='0F766E'

# Simple branded badge for cover (does not replace the official logo)
img=Image.new('RGB',(900,360),'white'); d=ImageDraw.Draw(img)
try:
    f1=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',94)
    f2=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',52)
    f3=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',30)
except:
    f1=f2=f3=None
# blue mark
d.rounded_rectangle((36,45,260,315),radius=42,fill='#0F172A')
d.ellipse((78,85,218,225),fill='#2563EB')
d.text((121,95),'E',font=f1,fill='white',anchor='ma')
d.text((305,92),'ERAGONOVA',font=f2,fill='#0F172A')
d.text((306,162),'ENTERPRISE',font=f3,fill='#2563EB')
d.text((306,222),'Technology • Engineering • Growth',font=f3,fill='#475569')
img.save(BADGE)

def font(run,name='Liberation Sans'):
    run.font.name=name
    rpr=run._element.get_or_add_rPr(); rf=rpr.rFonts
    if rf is None:
        rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    rf.set(qn('w:ascii'),name); rf.set(qn('w:hAnsi'),name)

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def borders(cell,color=BORDER,size='4'):
    tcPr=cell._tc.get_or_add_tcPr(); tb=tcPr.first_child_found_in('w:tcBorders')
    if tb is None: tb=OxmlElement('w:tcBorders'); tcPr.append(tb)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=tb.find(qn('w:'+edge))
        if el is None: el=OxmlElement('w:'+edge); tb.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),size); el.set(qn('w:color'),color); el.set(qn('w:space'),'0')

def style_runs(p,size=None,color=None,bold=None,name='Liberation Sans'):
    for r in p.runs:
        font(r,name)
        if size: r.font.size=Pt(size)
        if color: r.font.color.rgb=RGBColor.from_string(color)
        if bold is not None: r.bold=bold

def page_number(p):
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=p.add_run('Page '); font(r); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED)
    b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin')
    i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text='PAGE'
    e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end')
    r._r.extend([b,i,e])
    r2=p.add_run('  |  Eragonova Enterprise Website Owner’s Handbook'); font(r2); r2.font.size=Pt(8); r2.font.color.rgb=RGBColor.from_string(MUTED)

def heading(doc,text,level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.keep_with_next=True; return p

def body(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    r=p.add_run(text); font(r); return p

def bullets(doc,items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); r=p.add_run(x); font(r)

def steps(doc,items):
    for n,x in enumerate(items,1):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.55); p.paragraph_format.first_line_indent=Cm(-.55); p.paragraph_format.space_after=Pt(4)
        r=p.add_run(f'{n}. '); r.bold=True; font(r); r2=p.add_run(x); font(r2)

def note(doc,title,text,kind='info'):
    fills={'info':LIGHTBLUE,'safe':GREEN,'warn':AMBER,'danger':RED,'brand':'EEF2FF'}
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,fills.get(kind,LIGHTBLUE)); borders(c,BLUE if kind=='brand' else BORDER,'7' if kind=='brand' else '4')
    p=c.paragraphs[0]; r=p.add_run(title+'\n'); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); font(r)
    r2=p.add_run(text); r2.font.color.rgb=RGBColor.from_string(SLATE); font(r2); p.paragraph_format.space_after=Pt(0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def table(doc,headers,rows,widths=None,small=False):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False if widths else True
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=h; shade(c,NAVY); borders(c,'B8C4D4'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        style_runs(c.paragraphs[0],size=8.2 if small else 9,color=WHITE,bold=True)
        if widths: c.width=widths[j]
    for ri,row in enumerate(rows):
        cells=t.add_row().cells; fill=WHITE if ri%2==0 else LIGHT
        for j,val in enumerate(row):
            cells[j].text=str(val); shade(cells[j],fill); borders(cells[j],'D7E0EA'); cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            style_runs(cells[j].paragraphs[0],size=7.5 if small else 8.4)
            if widths: cells[j].width=widths[j]
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def code(doc,text):
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,'F1F5F9'); borders(c,'CBD5E1')
    p=c.paragraphs[0]; r=p.add_run(text); font(r,'Liberation Mono'); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(NAVY); p.paragraph_format.space_after=Pt(0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

# Document setup
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(1.65); sec.bottom_margin=Cm(1.55); sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7); sec.different_first_page_header_footer=True
styles=doc.styles; normal=styles['Normal']; normal.font.name='Liberation Sans'; normal.font.size=Pt(9.4); normal.font.color.rgb=RGBColor.from_string(NAVY); normal.paragraph_format.space_after=Pt(6)
for name,size,color,fname in [('Title',30,NAVY,'Liberation Serif'),('Heading 1',20,NAVY,'Liberation Serif'),('Heading 2',14,BLUE,'Liberation Sans'),('Heading 3',11,INDIGO,'Liberation Sans')]:
    st=styles[name]; st.font.name=fname; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True; st.paragraph_format.space_before=Pt(12 if name!='Heading 1' else 16); st.paragraph_format.space_after=Pt(6)
    if name=='Heading 1': st.paragraph_format.page_break_before=True
for s in doc.sections: page_number(s.footer.paragraphs[0])

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(26)
run=p.add_run(); run.add_picture(BADGE,width=Inches(5.3))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10)
r=p.add_run('WEBSITE OWNER’S HANDBOOK'); font(r,'Liberation Serif'); r.bold=True; r.font.size=Pt(26); r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('A plain-language guide to updating, maintaining and growing the Eragonova Enterprise website'); font(r); r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(MUTED)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12)
r=p.add_run('Practical Technology. Reliable Infrastructure. Sustainable Growth.'); font(r,'Liberation Serif'); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(BLUE)
cover=[('Live website','https://edward-wonder-mordey.github.io/eragonova-website/'),('GitHub repository','Edward-Wonder-Mordey/eragonova-website'),('Phone','+233 50 537 4537'),('Email','Eragonovaenterprise@gmail.com'),('Handbook version','Version 1.0 — 13 October 2026')]
t=doc.add_table(rows=len(cover),cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,(a,b) in enumerate(cover):
    t.cell(i,0).text=a; t.cell(i,1).text=b; shade(t.cell(i,0),'EFF6FF'); shade(t.cell(i,1),WHITE); borders(t.cell(i,0),'BFDBFE'); borders(t.cell(i,1),'BFDBFE'); style_runs(t.cell(i,0).paragraphs[0],size=8.3,bold=True,color=NAVY); style_runs(t.cell(i,1).paragraphs[0],size=8.3)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12)
r=p.add_run('Keep this document with the website. Update it whenever the site structure, services, contact details, policies or business systems change.'); font(r); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(MUTED)

# 1
heading(doc,'1. How to Use This Handbook',1)
body(doc,'This handbook is for both technical and non-technical readers. You do not need to understand React, JavaScript or web development to use most of it. Its main purpose is to tell you what controls each part of the website, where to make common changes, what is safe to edit, and when to get technical help.')
note(doc,'The safest rule','Change only the exact text, link, image name or list item you intend to change. Do not delete brackets, quotation marks, commas, import lines, function names or code structure unless you understand what they do.','warn')
heading(doc,'1.1 Three change levels',2)
table(doc,['Level','What it means','Examples'],[
('GREEN — Safe content change','Usually text, contact details, a service item, image or wording.','Phone number, email, service description, Academy track, partner wording.'),
('AMBER — Change carefully','Can affect navigation, layout, shared data or several pages.','Adding a page, changing routes, editing shared service data, changing SEO.'),
('RED — Technical/system change','Can affect deployment, security, forms, logins or future backend systems.','Formspree endpoint, deployment workflow, portal authentication, payments, secrets.')],widths=[Cm(3.4),Cm(6.2),Cm(7.8)])
heading(doc,'1.2 Before every website change',2)
steps(doc,['Open the live website and note exactly what you want to change.','Use Section 2 of this handbook to find the correct file.','Edit one logical item at a time in GitHub.','Use a clear commit message such as “Update Eragonova phone number” or “Add new Academy track”.','Wait for the GitHub Actions deployment to complete successfully.','Open the live website, refresh it, and test the changed page on desktop and mobile.','If something breaks, use GitHub file History to restore the previous working version.'])
note(doc,'What “commit” means','A commit is simply a saved version of the project. GitHub keeps the history so you can compare or restore earlier versions.','info')

# 2 quick finder
heading(doc,'2. Quick Finder — “I Want to Change This; Where Do I Go?”',1)
rows=[
('Phone / email / location','src/components/ContactSection.jsx','Edit the contacts list. Also check README if you want documentation to match.'),
('Contact form service choices','src/components/ContactSection.jsx','Edit serviceOptions.'),
('Where contact-form messages are sent','src/components/ContactSection.jsx','Form uses a Formspree endpoint. Change only if the Formspree account/form changes.'),
('Main menu','src/components/Header.jsx','Edit primaryItems or moreItems.'),
('Page routes','src/AppRouter.jsx','Maps a route such as services or academy to the page component.'),
('Homepage content','src/pages/HomePage.jsx','Controls most homepage sections and Elyra partnership panel.'),
('Homepage hero','src/components/HeroSection.new.jsx','This is the hero actually imported by HomePage.'),
('About page','src/pages/AboutPage.jsx','Mission, vision, values, positioning and operating model.'),
('Service divisions / service lists','src/data/siteData.js','Main shared source for the 16 divisions.'),
('Services page layout','src/pages/ServicesPage.jsx','Controls how pillars/divisions are presented.'),
('Projects page','src/pages/ClientFacingPages.jsx','Current ProjectsPage is exported from this file.'),
('Project cards on homepage/projects','src/components/ProjectsSection.new.jsx','Current project showcase component.'),
('Academy','src/pages/AcademyPage.jsx + src/data/siteData.js','Page layout in AcademyPage; track names in siteData.'),
('Resources / Community / Careers','src/pages/ClientFacingPages.jsx','All three are exported from this file.'),
('Elyra partnership','src/pages/PartnersPage.jsx and src/pages/HomePage.jsx','Logo, link and partnership wording appear in both.'),
('Client Portal','src/pages/PortalPage.jsx','Public portal presentation and access request experience.'),
('Contact page','src/pages/CorporatePages.jsx + ContactSection.jsx','Router uses ContactPage from CorporatePages; actual form/contact details are in ContactSection.'),
('Legal page','src/pages/CorporatePages.jsx','Router uses LegalPage from this file.'),
('Privacy policy','src/pages/PrivacyPage.jsx','Edit public privacy wording here.'),
('Terms','src/pages/TermsPage.jsx','Edit website/service terms here.'),
('Website colours and general styles','src/index.css','Global visual design and dark-mode styling.'),
('Capability/service visuals','src/components/CapabilityVisual.jsx and src/divisionVisual.css','Graphic previews used across service areas.'),
('Official logo / favicons','public/','eragonova.jpeg and browser icon files.'),
('SEO title/description/social preview','public/index.html','Main metadata for search and link sharing.'),
('Search engine sitemap','public/sitemap.xml','Update when major public routes change.'),
('Google Search Console verification','public/googlebcb43e0ea049d806.html','Do not remove unless changing verification method.'),
('Deployment',' .github/workflows/deploy-pages.yml','Builds the React site and publishes build/ to gh-pages.'),
]
table(doc,['What you want to update','Main file','Plain-language note'],rows,widths=[Cm(4.5),Cm(5.4),Cm(7.5)],small=True)

# 3 architecture
heading(doc,'3. Website Structure in Simple Terms',1)
body(doc,'The Eragonova website is a React website. Think of it as a set of page files, shared content files and reusable building blocks. GitHub stores the source. When you save a change to the main branch, GitHub Actions builds the website and publishes the finished build to the gh-pages branch.')
heading(doc,'3.1 The five parts to remember',2)
table(doc,['Part','What it does','Main location'],[
('Routes','Decides which page opens for #/about, #/services and other links.','src/AppRouter.jsx'),
('Pages','The main content for Home, About, Services, Academy, Partners and others.','src/pages/'),
('Shared content','Lists that feed multiple pages, especially the 16 divisions and sectors.','src/data/siteData.js'),
('Components','Reusable sections such as Header, Contact, Hero and project showcases.','src/components/'),
('Public files','Logo, browser icons, SEO page shell, sitemap and search-engine files.','public/')],widths=[Cm(3.1),Cm(7.7),Cm(6.6)])
heading(doc,'3.2 Current public routes',2)
table(doc,['Website address after #/','Page shown','Source file'],[
('home','Home','src/pages/HomePage.jsx'),('about','About','src/pages/AboutPage.jsx'),('services','Services','src/pages/ServicesPage.jsx'),('projects','Projects','src/pages/ClientFacingPages.jsx'),('academy','Academy','src/pages/AcademyPage.jsx'),('resources','Resources','src/pages/ClientFacingPages.jsx'),('community','Community','src/pages/ClientFacingPages.jsx'),('partners','Partners','src/pages/PartnersPage.jsx'),('careers','Careers','src/pages/ClientFacingPages.jsx'),('portal','Client Portal','src/pages/PortalPage.jsx'),('contact','Contact','src/pages/CorporatePages.jsx + ContactSection.jsx'),('legal','Legal','src/pages/CorporatePages.jsx'),('privacy','Privacy','src/pages/PrivacyPage.jsx'),('terms','Terms','src/pages/TermsPage.jsx')],widths=[Cm(3.7),Cm(4.2),Cm(9.5)],small=True)
note(doc,'Important legacy-file warning','CorporatePages.jsx still contains older copies of several pages, but AppRouter currently imports only ContactPage and LegalPage from it. Do not edit its older About/Services/Projects/Academy copies expecting the live pages to change. Use the files listed above.','danger')

# 4 brand/contact
heading(doc,'4. Brand Identity and Contact Information',1)
body(doc,'Eragonova is positioned publicly as a multidisciplinary technology, engineering and business solutions company. The public tagline is “Technology • Engineering • Growth,” and the core direction line is “Practical Technology. Reliable Infrastructure. Sustainable Growth.”')
table(doc,['Item','Current public value','Where to update'],[
('Business name','Eragonova Enterprise','Several page headings; Header.jsx; public/index.html'),
('Tagline','Technology • Engineering • Growth','public/index.html and brand copy where used'),
('Direction line','Practical Technology. Reliable Infrastructure. Sustainable Growth.','HomePage.jsx / hero-related copy'),
('Phone','+233 50 537 4537','ContactSection.jsx'),
('Email','Eragonovaenterprise@gmail.com','ContactSection.jsx'),
('Location','Kumasi, Ghana','ContactSection.jsx'),
('Live URL','https://edward-wonder-mordey.github.io/eragonova-website/','package.json, public/index.html, sitemap/robots as applicable'),
('Official logo','public/eragonova.jpeg','Replace carefully; keep filename if you want existing references to keep working')],widths=[Cm(3.7),Cm(6.3),Cm(7.4)])
heading(doc,'4.1 Changing contact details safely',2)
steps(doc,['Open src/components/ContactSection.jsx.','Find the contacts array near the top/middle of the file.','Change the visible value and the matching tel: or mailto: link if applicable.','If the phone changes, change both the displayed number and the tel:+233... link.','Commit the change and test the Contact page after deployment.'])

# 5 file map
heading(doc,'5. File and Folder Map',1)
rows=[
('src/AppRouter.jsx','Route map','AMBER'),('src/pages/HomePage.jsx','Homepage','GREEN/AMBER'),('src/pages/AboutPage.jsx','About','GREEN'),('src/pages/ServicesPage.jsx','Services presentation','AMBER'),('src/pages/AcademyPage.jsx','Academy','GREEN/AMBER'),('src/pages/ClientFacingPages.jsx','Projects, Resources, Community, Careers','GREEN/AMBER'),('src/pages/PartnersPage.jsx','Strategic Partners / Elyra Suite','GREEN'),('src/pages/PortalPage.jsx','Client Portal public experience','AMBER'),('src/pages/PrivacyPage.jsx','Privacy policy','GREEN/AMBER'),('src/pages/TermsPage.jsx','Terms','GREEN/AMBER'),('src/pages/CorporatePages.jsx','Current Contact + Legal; also contains older duplicate page code','CAUTION'),('src/data/siteData.js','16 divisions, pillars, sectors, Academy tracks, resources and community programmes','GREEN/AMBER'),('src/components/Header.jsx','Navigation + dark-mode toggle','AMBER'),('src/components/ContactSection.jsx','Contact details + Formspree form','AMBER'),('src/components/HeroSection.new.jsx','Current homepage hero','GREEN/AMBER'),('src/components/ProjectsSection.new.jsx','Current project showcase','GREEN/AMBER'),('src/components/CapabilityVisual.jsx','Reusable service/capability graphics','AMBER'),('src/index.css','Global styling','AMBER/RED'),('public/index.html','SEO/social preview shell','AMBER'),('public/eragonova.jpeg','Official logo asset','GREEN'),('public/sitemap.xml / robots.txt','Search-engine discovery files','AMBER'),('.github/workflows/deploy-pages.yml','Automatic build/deployment','RED'),('package.json','React dependencies, scripts and homepage URL','RED')]
table(doc,['File / folder','Purpose','Change level'],rows,widths=[Cm(6.2),Cm(8.5),Cm(2.7)],small=True)

# 6 home/about
heading(doc,'6. Updating the Home and About Pages',1)
heading(doc,'6.1 Homepage',2)
body(doc,'The homepage is mainly in src/pages/HomePage.jsx. Its first hero section comes from src/components/HeroSection.new.jsx. The homepage then uses shared service pillars, sectors and Academy tracks from siteData.js, a current project showcase from ProjectsSection.new.jsx, plus the Elyra strategic-partnership block.')
bullets(doc,['Hero wording: HeroSection.new.jsx','Homepage section wording, partnership and call-to-actions: HomePage.jsx','Service pillars and sectors: siteData.js','Projects shown: ProjectsSection.new.jsx','Academy track names: siteData.js'])
heading(doc,'6.2 About page',2)
body(doc,'Use src/pages/AboutPage.jsx for the live About page. This is where mission, vision, values, company positioning and operating model should be updated.')
note(doc,'Do not use the duplicate AboutPage inside CorporatePages.jsx','That older copy is not the live About route. AppRouter imports src/pages/AboutPage.jsx directly.','warn')

# 7 services
heading(doc,'7. Updating Services and the 16 Divisions',1)
body(doc,'The most important service information lives in src/data/siteData.js. This shared data feeds the Services page and other parts of the site. If a service name or item changes here, it can update more than one page automatically.')
heading(doc,'7.1 Current service divisions',2)
services=[
('01','IT & Systems Administration'),('02','Networking & Infrastructure'),('03','Cloud & Digital Transformation'),('04','Data & Analytics'),('05','Cybersecurity & Information Protection'),('06','Software & Digital Solutions'),('07','Electrical & Electronic Engineering'),('08','Administration & Business Support'),('09','Documentation & Compliance Services'),('10','Branding, Media & Print Services'),('11','Research & Intelligence Services'),('12','Organisational Development'),('13','Consulting Services'),('14','Managed Services'),('15','Eragonova Academy'),('16','Community Development Programmes')]
table(doc,['No.','Division','Main edit location'],[(a,b,'src/data/siteData.js') for a,b in services],widths=[Cm(1.5),Cm(10),Cm(6)],small=True)
heading(doc,'7.2 Add a service under an existing division',2)
steps(doc,['Open src/data/siteData.js.','Find the division title.','Inside its groups list, find the correct group such as End User Support, Network Design or Corporate Documentation.','Add the new service as another quoted item, keeping commas and brackets consistent.','Commit and check the Services page.'])
heading(doc,'7.3 Add a completely new division',2)
body(doc,'This is an AMBER-level change because the new division also needs a number, icon, summary, groups, and a place inside servicePillars. If a 17th division is added, also review ServicesPage.jsx, any visual grouping logic and the contact-form service list.')

# 8 client-facing content
heading(doc,'8. Projects, Resources, Community and Careers',1)
body(doc,'These four public pages currently come from one file: src/pages/ClientFacingPages.jsx. It exports separate page functions for each route.')
table(doc,['Page','Where content is controlled','Related shared data/components'],[
('Projects','ClientFacingPages.jsx','ProjectsSection.new.jsx'),('Resources','ClientFacingPages.jsx','resourceTypes in siteData.js'),('Community','ClientFacingPages.jsx','communityPrograms in siteData.js'),('Careers','ClientFacingPages.jsx','Mostly page-local content and contact links')],widths=[Cm(3.2),Cm(6.5),Cm(7.8)])
heading(doc,'8.1 Adding real case studies later',2)
body(doc,'Use genuine project information only. A strong case study should show the challenge, approach, deliverables and outcome. Do not invent client names, results or percentages. If a client must remain confidential, describe the sector and scope without identifying them.')

# 9 academy
heading(doc,'9. Eragonova Academy',1)
body(doc,'The Academy page layout is in src/pages/AcademyPage.jsx. The shared Academy track names and subjects are in src/data/siteData.js under academyTracks and Division 15.')
table(doc,['Training area','Current examples'],[
('Technology','Networking, Linux, Windows Server, Active Directory, Cloud Computing'),('Data','Excel, SQL, Power BI, Tableau'),('Business','Administrative Support, Customer Service, Business Writing, Documentation'),('Professional Development','Leadership, Communication, Project Coordination'),('Engineering','Electrical Fundamentals, Electronics, Technical Safety')],widths=[Cm(5.2),Cm(12.3)])
note(doc,'Certification wording','Only describe a course as accredited, certified by a named body or officially recognised if that is true and documented. Otherwise use wording such as training, programme, workshop or completion certificate.','warn')

# 10 partnership
heading(doc,'10. Strategic Partnership with Elyra Suite',1)
body(doc,'Elyra Suite is shown as an independent strategic partner, not a division of Eragonova. The reciprocal partnership appears on the Eragonova homepage and on the dedicated Partners page.')
table(doc,['Where','File','What can be updated'],[
('Homepage partnership panel','src/pages/HomePage.jsx','Elyra logo URL, Elyra website URL, tagline and summary'),('Dedicated Partners page','src/pages/PartnersPage.jsx','ELYRA object, partnership wording, principles and link')],widths=[Cm(4.5),Cm(6.3),Cm(6.7)])
note(doc,'Keep independence clear','Do not change the wording to imply Elyra Suite is owned by, merged into, or a department of Eragonova unless the real business relationship changes.','brand')

# 11 portal
heading(doc,'11. Client Portal',1)
body(doc,'The public Client Portal experience is in src/pages/PortalPage.jsx. Treat the current portal as a controlled access / engagement experience. Do not claim there is a full authentication, billing or database backend unless one is actually connected.')
heading(doc,'11.1 Future portal upgrades',2)
bullets(doc,['Real user authentication and password reset','Client-specific invoices and documents','Secure training materials','Appointment/project status','Payment records','Role-based access and audit logging'])
note(doc,'Security rule','Never place client passwords, API keys, payment secrets or private documents directly inside public React source files or the public GitHub repository.','danger')

# 12 contact/formspree
heading(doc,'12. Contact Form and Formspree',1)
body(doc,'The visible contact details and the contact form are both in src/components/ContactSection.jsx. The form currently sends submissions to Formspree using the endpoint https://formspree.io/f/mwvegdqp.')
heading(doc,'12.1 If contact details change',2)
body(doc,'Edit the contacts array in ContactSection.jsx. Keep the displayed phone/email and the clickable link values consistent.')
heading(doc,'12.2 If the Formspree account or form changes',2)
steps(doc,['Create or confirm the replacement Formspree form.','Copy the new Formspree endpoint.','Open ContactSection.jsx.','Find fetch("https://formspree.io/f/mwvegdqp", ...).','Replace only the URL with the new endpoint.','Deploy and submit a test message from the live website.','Confirm the test appears in the intended Formspree account/email destination.'])
note(doc,'Do not publish private keys','The current Formspree URL is a public form endpoint. Secret API keys, passwords or tokens should never be added to this frontend file.','danger')

# 13 legal
heading(doc,'13. Legal, Privacy and Terms Pages',1)
table(doc,['Page','Current file','Update when'],[
('Legal','src/pages/CorporatePages.jsx','Business/legal notices or legal-page content changes.'),('Privacy','src/pages/PrivacyPage.jsx','Data collection, forms, analytics, cookies, portal or new external services change.'),('Terms','src/pages/TermsPage.jsx','Service terms, website terms or operational terms change.')],widths=[Cm(3.2),Cm(6.3),Cm(8.0)])
body(doc,'For legal wording, use the website text as a practical business policy, not a substitute for professional legal advice where a formal contract, regulated service or high-value transaction requires it.')

# 14 nav/routing/theme
heading(doc,'14. Navigation, Routes and Dark Mode',1)
heading(doc,'14.1 Navigation menu',2)
body(doc,'The menu is controlled by src/components/Header.jsx. primaryItems are shown directly on large screens. moreItems appear inside the More menu. On mobile, both lists are combined into the mobile menu.')
heading(doc,'14.2 Adding a page to navigation',2)
steps(doc,['Make sure the page component exists.','Add the route to the routes object in src/AppRouter.jsx.','Add a matching menu item in Header.jsx.','Use the same route word in both places.','Commit and test the link on desktop and mobile.'])
heading(doc,'14.3 Dark mode',2)
body(doc,'The Header includes a sun/moon toggle. The selected theme is saved in the browser using localStorage under the name eragonova-theme. Dark-mode colours are mainly defined through Tailwind classes and src/index.css.')

# 15 assets/visuals
heading(doc,'15. Images, Logos and Visual Assets',1)
body(doc,'Files inside public/ are copied into the built website as public assets. The official Eragonova logo is public/eragonova.jpeg. Browser icons and manifest files are also in public/.')
table(doc,['Asset','Location','Notes'],[
('Main Eragonova logo','public/eragonova.jpeg','Header and social preview references rely on this file.'),('Browser icons','public/favicon.ico and PNG icon files','Used in browser tabs, bookmarks and devices.'),('Social sharing image','public/index.html points to eragonova.jpeg','Change og:image/twitter:image only if you add a better sharing image.'),('Elyra logo','Remote URL in HomePage.jsx and PartnersPage.jsx','Comes from Elyra Suite repository.'),('Capability visuals','CapabilityVisual.jsx + divisionVisual.css','Generated graphic-style UI previews, not photographs.')],widths=[Cm(4),Cm(6.5),Cm(7.0)])
heading(doc,'15.1 Replacing the main logo',2)
steps(doc,['Prepare the approved logo image.','Upload it to public/.','If you keep the filename eragonova.jpeg, existing references continue to work.','If you change the filename, update Header.jsx and public/index.html references.','Test header, browser/social preview metadata and mobile layout.'])

# 16 styling
heading(doc,'16. Colours, Layout and Styling',1)
body(doc,'The site uses a dark slate and blue/indigo visual system with accent colours for different capability areas. Global CSS lives in src/index.css. Some specialised visual styling also lives in src/divisionVisual.css. Tailwind utility classes are written directly inside JSX files.')
note(doc,'Styling caution','Changing one CSS class can affect many pages. Save a copy or rely on GitHub history, change one thing at a time, and test light mode and dark mode.','warn')
heading(doc,'16.1 Files you normally do not need to touch',2)
bullets(doc,['tailwind.config and tailwind.config.js — framework configuration','postcss.config — CSS build configuration','src/index.js — application startup','node_modules — installed packages; do not manually edit','build — generated site output; normally let the deployment process recreate it'])

# 17 SEO/search
heading(doc,'17. SEO, Search Engines and Social Sharing',1)
body(doc,'The main SEO description, title, keywords, canonical URL and social-sharing metadata are in public/index.html. Search-engine support files are in public/robots.txt and public/sitemap.xml. A Google Search Console verification file is already present.')
table(doc,['Need','File','What to edit'],[
('Change Google/search description','public/index.html','meta name="description"'),('Change browser/search title','public/index.html','<title> and og:title/twitter:title'),('Change social preview description','public/index.html','og:description and twitter:description'),('Change social preview image','public/index.html','og:image and twitter:image'),('Update indexed routes','public/sitemap.xml','Add/remove public URLs when site structure changes'),('Control crawling','public/robots.txt','Usually leave as-is unless SEO strategy changes'),('Google verification','public/googlebcb43e0ea049d806.html','Do not remove while using this verification method')],widths=[Cm(4.2),Cm(5.7),Cm(7.6)],small=True)
heading(doc,'17.1 Analytics',2)
body(doc,'No Google Analytics gtag integration is currently visible in the source reviewed for this handbook. If analytics is added later, document the measurement ID location here and update the Privacy page if the tool uses cookies or collects visitor data.')

# 18 GitHub/deploy
heading(doc,'18. Editing in GitHub and Publishing Changes',1)
heading(doc,'18.1 Simple GitHub editing',2)
steps(doc,['Open the eragonova-website repository.','Open the correct file from Section 2.','Click the pencil/edit icon.','Make the smallest necessary change.','Use Preview changes or review the diff.','Commit directly to main only when you are comfortable with the change.','Open the Actions tab and wait for Build and Deploy Eragonova Website to finish.','Open the live site and test the change.'])
heading(doc,'18.2 What the deployment does',2)
body(doc,'The workflow in .github/workflows/deploy-pages.yml runs whenever main changes. It checks out the repository, sets up Node.js, installs packages, runs npm run build, then publishes the build folder to the gh-pages branch using peaceiris/actions-gh-pages.')
code(doc,'main branch change  →  npm install  →  npm run build  →  build/  →  gh-pages branch  →  live GitHub Pages website')
note(doc,'Do not edit gh-pages as the main source','The gh-pages branch is deployment output. Make normal website changes on main so the next build does not overwrite manual edits.','danger')

# 19 troubleshooting
heading(doc,'19. Troubleshooting Common Problems',1)
table(doc,['Problem','Likely reason','What to do'],[
('I changed a file but nothing changed live','Wrong file or deployment still running','Check the route/file map, then GitHub Actions.'),('Old version still appears','Browser cache','Hard refresh with Ctrl+F5 or test in private/incognito window.'),('Site build failed','Code syntax, missing comma/bracket, dependency issue','Open Actions logs; restore last good edit if necessary.'),('A menu link shows 404','Route missing or route word mismatch','Check AppRouter.jsx and Header.jsx use the same route.'),('Contact form does not send','Formspree/network/endpoint issue','Test Formspree endpoint and check ContactSection.jsx.'),('Dark mode looks wrong','A component lacks dark: classes or global CSS conflict','Check component classes and index.css.'),('Elyra logo does not load','Remote logo URL changed or blocked','Check ELYRA.logo in HomePage.jsx and PartnersPage.jsx.'),('Service changed on one page but not another','Edited a page copy instead of shared data','Update siteData.js if the content is shared.'),('About/Services edit in CorporatePages had no effect','That is an older duplicate','Use AboutPage.jsx / ServicesPage.jsx listed in this handbook.')],widths=[Cm(4.2),Cm(6.1),Cm(7.2)],small=True)

# 20 security/backends
heading(doc,'20. Security, Client Data and Future Backend Features',1)
bullets(doc,['Never commit passwords, private API keys, payment secrets or client confidential documents to the public repository.','Do not store real portal usernames/passwords in React source code.','Use a secure backend or trusted service for authentication, payments and private client files.','Limit Formspree access to the appropriate business owners and review spam/security settings.','Use least-privilege access when adding collaborators to GitHub.','Before publishing project/client information, confirm permission and remove confidential details.'])
heading(doc,'20.1 Features that need real backend systems',2)
table(doc,['Feature','What is needed'],[
('Authenticated Client Portal','Identity/authentication service, secure database/storage, access rules, password recovery.'),('Online payments','Payment provider integration, server-side verification/webhooks, terms/refund process.'),('Private documents','Secure storage with client-specific permissions.'),('Real Academy accounts','Learner database, enrolment records, protected materials and completion records.'),('Automated invoicing','Accounting/payment workflow or secure backend/API.')],widths=[Cm(5),Cm(12.5)])

# 21 maintenance
heading(doc,'21. Maintenance Checklist',1)
heading(doc,'Monthly',2)
bullets(doc,['Open every main navigation link and check for broken pages.','Submit a test contact form and confirm it arrives.','Check phone/email links.','Review project, Academy, resource and careers wording for outdated information.','Check the Elyra partner link/logo.','Test light and dark modes on desktop and mobile.','Check GitHub Actions for repeated failures.'])
heading(doc,'Quarterly',2)
bullets(doc,['Review all 16 divisions against what Eragonova actually offers.','Update real project/case-study evidence.','Review Privacy, Terms and Legal pages if tools or services changed.','Review public/index.html description and social preview.','Check sitemap and Search Console indexing.','Remove genuinely unused files only after confirming they are not imported anywhere.','Update this handbook and its change log.'])
heading(doc,'Before a major launch or proposal',2)
bullets(doc,['Confirm contact details and business positioning.','Check project examples are accurate and permission-safe.','Confirm services in proposals match the website.','Test the contact form from another device/network.','Check mobile navigation and all CTA buttons.'])

# 22 growth roadmap
heading(doc,'22. Recommended Future Enhancements',1)
table(doc,['Enhancement','Why it helps','Do when'],[
('Real project case studies','Builds proof and trust using genuine outcomes.','As verified client/project work becomes available.'),('Analytics','Shows traffic, popular pages and enquiry paths.','When you are ready to manage analytics/privacy.'),('Custom domain','Stronger corporate identity than the GitHub Pages URL.','When domain ownership is ready.'),('Professional domain email','Improves brand consistency.','After domain setup.'),('Portal authentication','Turns public portal concept into real client access.','When recurring clients justify it.'),('Resource library','Makes templates/guides genuinely downloadable and useful.','As original resources are produced.'),('Academy registration system','Supports cohorts, schedules and learner records.','When training volume grows.'),('Structured SEO per route','Improves search targeting beyond one HTML metadata shell.','When SEO becomes a dedicated growth channel.')],widths=[Cm(4.4),Cm(7.6),Cm(5.5)],small=True)

# Appendix A finer services
heading(doc,'Appendix A. Current Service Catalogue Reference',1)
service_groups=[
('01 IT & Systems Administration','Computer troubleshooting; printer/software/help desk; Windows/Linux Server; Active Directory; Group Policy; storage; monitoring; backup; disaster recovery; virtualisation; Microsoft 365; Azure; AWS; Google Workspace.'),
('02 Networking & Infrastructure','LAN/WAN/wireless/campus design; DNS; DHCP; routing; switching; structured cabling; racks; performance reviews; infrastructure assessments; upgrade recommendations.'),
('03 Cloud & Digital Transformation','Azure/AWS/hybrid cloud; process mapping; workflow optimisation; migration; Infrastructure as Code; Terraform; DevOps; containers; Kubernetes; cloud security.'),
('04 Data & Analytics','SQL reporting; validation/cleaning/auditing; Power BI; Tableau; dashboards; KPIs; operational reporting; decision support; predictive analytics; AI-assisted reporting; governance.'),
('05 Cybersecurity & Information Protection','Security assessments; vulnerability reviews; risk analysis; log reviews; monitoring; documentation; staff awareness; cyber hygiene; policy development; authorised response support.'),
('06 Software & Digital Solutions','Corporate/portfolio/e-commerce sites; booking systems; client portals; learning platforms; mobile apps; SaaS products; internal systems.'),
('07 Electrical & Electronic Engineering','Wiring design; electrical installation; power distribution; electronics diagnostics/repairs/circuit analysis; drawings/reports/compliance; solar; smart buildings; CCTV; access control.'),
('08 Administration & Business Support','Virtual/executive assistance; calendars; email; meetings; records; data entry; customer support; process administration; workflow documentation; administrative audits.'),
('09 Documentation & Compliance','Company profiles; business plans; reports; presentations; tenders; technical proposals; bid packaging; SOPs; policies; procedures; manuals; user guides; training materials; knowledge bases.'),
('10 Branding, Media & Print','Logo design; identity systems; visual guides; marketing/advertising design; corporate/training/large-format printing; T-shirts; gifts; merchandise.'),
('11 Research & Intelligence','Market research; community assessments; feasibility; survey analysis; programme evaluation; impact assessment; data collection; reporting; decision support.'),
('12 Organisational Development','Strategic plans; roadmaps; process reviews; quality assurance; operational assessments; stakeholder management; reporting systems; performance reviews.'),
('13 Consulting Services','Technology, engineering, operations, data, administrative and organisational consulting.'),
('14 Managed Services','Remote IT support; monitoring; maintenance; managed reports/policies/knowledge; monthly reporting; dashboard maintenance.'),
('15 Eragonova Academy','Technology, data, business, professional-development and engineering training.'),
('16 Community Development Programmes','Digital literacy; youth technology; entrepreneurship; community projects; church technology support; self-reliance programmes.')]
for title,desc in service_groups:
    heading(doc,title,2); body(doc,desc)

# Appendix B active/legacy
heading(doc,'Appendix B. Active Files vs Older/Duplicate Files',1)
table(doc,['File','Current status','Owner guidance'],[
('HeroSection.new.jsx','ACTIVE — imported by HomePage','Edit this for the live homepage hero.'),('HeroSection.jsx','Older/alternate file','Do not assume changes show live unless import is changed.'),('ProjectsSection.new.jsx','ACTIVE — imported by HomePage and ClientFacingPages','Edit current project showcase here.'),('ProjectsSection.jsx','Older/alternate file','Do not use for live changes unless imports change.'),('ServicesSection.new.jsx / ServicesSection.jsx','Reusable/older components; not the main Services route','Current Services route is src/pages/ServicesPage.jsx.'),('CorporatePages.jsx','MIXED','ContactPage and LegalPage are routed. Older duplicate page exports in this file are not the live About/Services/Projects/Academy routes.'),('build/','Generated output','Do not treat as the master source.'),('gh-pages branch','Published output','Do not use as normal editing branch.'),('node_modules/','Installed packages','Do not manually edit.')],widths=[Cm(5.3),Cm(5.5),Cm(6.7)],small=True)

# Appendix C change log
heading(doc,'Appendix C. Website Change Log Template',1)
body(doc,'Record important changes here or maintain the same table in a separate document. This makes it easier to understand what changed and why.')
table(doc,['Date','Changed by','Area / file','What changed','Reason / notes'],[
('13 Oct 2026','Initial handbook','docs/','Created website owner’s handbook.','Baseline documentation.'),('','','','',''),('','','','',''),('','','','',''),('','','','',''),('','','','','')],widths=[Cm(2.5),Cm(3),Cm(3.5),Cm(4.5),Cm(4.0)],small=True)

heading(doc,'Appendix D. Final Owner Checklist',1)
bullets(doc,['I know the main source branch is main.','I know gh-pages is generated deployment output.','I know the correct files for Home, About, Services, Academy, Partners and Portal.','I know shared service content belongs mainly in siteData.js.','I know contact details and the Formspree endpoint are in ContactSection.jsx.','I know CorporatePages.jsx contains both active Contact/Legal code and older duplicate code.','I know public/index.html controls the main SEO/social metadata.','I know .github/workflows/deploy-pages.yml controls automatic deployment.','I will not store passwords, secrets or client confidential data in the public repository.','After major site changes, I will update this handbook too.'])
note(doc,'Handbook maintenance rule','Whenever the website structure, routes, contact system, services, portal, deployment or business model changes, update this handbook in the same period so it remains trustworthy.','brand')

# metadata
props=doc.core_properties; props.title='Eragonova Enterprise Website Owner’s Handbook'; props.subject='Plain-language website ownership, update and maintenance guide'; props.author='Eragonova Enterprise'; props.keywords='Eragonova, website handbook, GitHub, maintenance, owner guide'
doc.save(OUT)
print(OUT)
